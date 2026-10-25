import docx
from typing import List, Tuple
from .base import BaseExtractor


class DocxExtractor(BaseExtractor):
    """Extracts text from DOCX files by paragraphs and tables."""

    def extract_pages(self, file_path: str) -> List[Tuple[int, str]]:
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n\n".join(paragraphs).strip()
            if not full_text:
                raise ValueError("The Word document is empty or contains no readable text.")

            # DOCX has no explicit native page boundaries without layout engine, so assign page 1
            return [(1, full_text)]
        except Exception as e:
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
